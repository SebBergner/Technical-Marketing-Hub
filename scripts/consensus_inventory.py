#!/usr/bin/env python
"""Consensus videos against SharePoint: the inventory to review before migrating.

READ-ONLY. Nothing is written to Consensus or SharePoint. It calls:

    POST /api/reports/v1.0/demosDetails   every published demo, with its video
                                          uuid, exact duration and playlist chapters
    POST /api/integr/v1.0/demo/search     every demo: language, folder, and the
                                          unpublished ones

V2 (`/api/v2/demos/search`) is deliberately not used: its demo uuids are a
different scheme (1 of 661 joined, 2026-10-09) and its `usage` was 0 on all
but 2 demos.

and reads the SharePoint side from this machine's mirror (DATA_DIR/mirror:
Demo Video and Demo Catalog). Sync first if the mirror is old; the summary
sheet says how old it is.

The matching itself is backend/services/consensus_inventory.py.

Usage, from the repo root:

    python scripts/consensus_inventory.py
    python scripts/consensus_inventory.py --out "C:/Work/TDD Hub/Highlevel/inventory.xlsx"
    python scripts/consensus_inventory.py --details saved.json   # reuse a saved fetch
"""
from __future__ import annotations

import argparse
import io
import json
import os
import sys
from datetime import date, datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from openpyxl import Workbook                                   # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill        # noqa: E402
from openpyxl.utils import get_column_letter                    # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation    # noqa: E402

from backend.config import settings                             # noqa: E402
from backend.services import consensus_inventory as ci          # noqa: E402

DETAILS_PAGE = 50          # the reports endpoint refuses 100 (measured 2026-10-09)

FILLS = {
    ci.IN_SP: "D9EAD3", ci.LIKELY: "FFF2CC", ci.OTHER_LANGUAGE: "DDEBF7", ci.OTHER_CUT: "FCE5CD",
    ci.NO_DURATION: "FCE5CD", ci.MIGRATE: "F4CCCC",
}
DECISIONS = ["Same video", "Not the same video", "Migrate", "Do not migrate"]


# ───────────────────────────────────────────────────────────────── fetching
def fetch_details() -> list[dict]:
    from backend.integrations.consensus import PATHS, get_consensus_client
    client = get_consensus_client()
    items, page = [], 1
    while True:
        data = client._post(PATHS["demos_details"],
                            {"paging": {"limit": DETAILS_PAGE, "page": page}})["data"]
        items += data["items"]
        nxt = (data.get("paging") or {}).get("nextPage")
        if not nxt or nxt == page:
            return items
        page = nxt


def fetch_all_v1() -> list[dict]:
    from backend.integrations.consensus import get_consensus_client
    demos = get_consensus_client().list_demos(limit=5000, include_archived=True,
                                              published_only=False)
    return [d.raw for d in demos]


def mirror(name: str) -> tuple[list[dict], str | None]:
    path = os.path.join(settings.data_dir, "mirror", f"{name}.json")
    if not os.path.exists(path):
        return [], None
    with open(path, encoding="utf-8") as fh:
        rows = json.load(fh)
    return rows, datetime.fromtimestamp(os.path.getmtime(path)).strftime("%Y-%m-%d %H:%M")


# ───────────────────────────────────────────────────────────────── writing
def mmss(seconds) -> str:
    if seconds is None:
        return ""
    seconds = int(round(seconds))
    return f"{seconds // 60}:{seconds % 60:02d}"


#: Columns whose values are links, made clickable.
LINK_COLUMNS = {"Open in Consensus", "SharePoint link"}


def sheet(wb, title: str, headers: list[str], rows: list[list], widths: dict | None = None,
          status_col: int | None = None, decision_col: int | None = None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5F3A")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, 16)
        if h in LINK_COLUMNS:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("http"):
                        cell.hyperlink = cell.value
                        cell.style = "Hyperlink"
    if status_col:
        for row in ws.iter_rows(min_row=2, min_col=status_col, max_col=status_col):
            for cell in row:
                if cell.value in FILLS:
                    cell.fill = PatternFill("solid", fgColor=FILLS[cell.value])
    if decision_col and rows:
        dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"',
                            allow_blank=True)
        ws.add_data_validation(dv)
        letter = get_column_letter(decision_col)
        dv.add(f"{letter}2:{letter}{len(rows) + 1}")
    return ws


def write(out_path: str, result: dict, v1_all: list[dict], sp_age: dict,
          fetched_at: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    videos = result["videos"]
    counts = {s: sum(1 for _, m in videos if m.status == s) for s in ci.STATUS_ORDER}
    published = {d["uuid"] for d in result["details"]}
    unpublished = [d for d in v1_all if d.get("uuid") not in published]
    summary = [
        ["Consensus videos against SharePoint — inventory for review", ""],
        ["Fetched from Consensus", fetched_at],
        ["SharePoint mirror (Demo Video)", sp_age.get("demo_video") or "missing"],
        ["SharePoint mirror (Demo Catalog)", sp_age.get("sharepoint") or "missing"],
        ["", ""],
        ["Published Consensus demos (reports endpoint)", len(result["details"])],
        ["  of which playlists (standard / advanced / flow)", len(result["playlists"])],
        ["  single demos without a video", len(result["no_video"])],
        ["Unpublished or draft demos (not in this inventory)", len(unpublished)],
        ["Distinct Consensus videos", len(videos)],
        ["  used only as a playlist chapter (no demo of its own)",
         sum(1 for v, _ in videos if not v.demos)],
        ["SharePoint video files compared", len(result["sharepoint"])],
        ["", ""],
    ] + [[s, counts[s]] for s in ci.STATUS_ORDER] + [
        ["", ""],
        ["How to read it", ""],
        ["One row per Consensus video (by video uuid): a video can be a demo of its own and a "
         "chapter of several playlists; the Videos sheet lists both.", ""],
        ["To find one in Consensus: 'Where in Consensus' says which demo or playlist it is "
         "in, and 'Open in Consensus' opens that page. A video with no demo of its own is a "
         "chapter: open the playlist and pick the chapter. The video uuid is internal to "
         "Consensus and cannot be searched for there.", ""],
        ["A match needs the name AND the length to agree. Length alone is not evidence: a "
         "video typically has ~7 unrelated SharePoint videos within 2 seconds.", ""],
        ["File size is not available from Consensus, so it cannot be compared. It is shown "
         "for SharePoint so copies can be told apart.", ""],
        ["Review: set 'Review decision' on each row that says 'check'. 'SharePoint item id' is "
         "where the Consensus UUID would be written (column ConsensusUUID).", ""],
        ["Open question: ConsensusUUID holds the single demo's uuid. A video used only as a "
         "playlist chapter has no demo uuid — its video uuid is listed instead.", ""],
    ]
    ws = wb.create_sheet("Summary")
    for r in summary:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 95
    ws.column_dimensions["B"].width = 22
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        if row[0].value in FILLS:
            row[0].fill = PatternFill("solid", fgColor=FILLS[row[0].value])

    headers = ["Status", "Review decision", "Consensus title", "Where in Consensus",
               "Open in Consensus", "Consensus demo UUID",
               "Video UUID", "Length", "Language", "Public", "Playlists using it",
               "Consensus folder", "Created",
               "SharePoint file", "Library", "SharePoint demo", "SP length", "Length diff (s)",
               "Size (MB)", "Resolution", "Name score", "Copies of this file",
               "Next candidate", "SharePoint item id", "SharePoint link", "Notes"]
    v1 = {d.get("uuid"): d for d in v1_all}

    def link(uuid: str) -> str:
        """The demo's or playlist's Consensus page. The video uuid itself is
        Consensus-internal: nothing in its interface finds it."""
        return (v1.get(uuid) or {}).get("previewLink") or             f"https://play.goconsensus.com/{uuid}?preview=sales"

    def private(public: bool) -> str:
        return "" if public else " (not public)"

    def where(v) -> tuple[str, str]:
        """How a person finds this video in Consensus, and the page to open."""
        if v.demos:
            d = v.demos[0]
            text = f'Demo "{d["title"]}"{private(d["public"])}'
            if v.playlists:
                text += f"; also a chapter of {len(v.playlists)} playlist(s)"
            return text, link(d["uuid"])
        p = v.playlists[0]
        text = f'Chapter "{p["chapter"]}" of playlist "{p["title"]}"{private(p["public"])}'
        if len(v.playlists) > 1:
            text += f" and {len(v.playlists) - 1} more"
        return text, link(p["uuid"])

    rows = []
    for v, m in videos:
        demo = v.demos[0] if v.demos else {}
        folder = ((v1.get(demo.get("uuid")) or {}).get("folderInfo") or {}).get("name") or ""
        b = m.best or {}
        rows.append([
            m.status, "", v.title, *where(v),
            ", ".join(d["uuid"] for d in v.demos), v.video_uuid,
            mmss(v.duration), v.language or "", "yes" if v.public else "no",
            "; ".join(f'{p["title"]} › {p["chapter"]}' for p in v.playlists), folder,
            (demo.get("created_at") or "")[:10],
            b.get("file"), b.get("library"), b.get("asset_title"), mmss(b.get("duration")),
            m.duration_diff, round(b["size_bytes"] / 1048576, 1) if b.get("size_bytes") else None,
            f'{b["width"]}x{b["height"]}' if b.get("width") else "",
            m.name_score or None,
            "; ".join(f'{c["library"]}: {c["asset_title"]} / {c["file"]}' for c in m.copies),
            f'{m.runner_up["file"]} ({m.runner_up_score})' if m.runner_up else "",
            b.get("item_id"), b.get("web_url"), ""])
    widths = {"Status": 30, "Review decision": 18, "Consensus title": 45,
              "Where in Consensus": 55, "Open in Consensus": 30,
              "Consensus demo UUID": 38, "Video UUID": 38, "Playlists using it": 45,
              "Consensus folder": 32, "SharePoint file": 48,
              "SharePoint demo": 38, "Copies of this file": 40, "Next candidate": 45,
              "SharePoint item id": 36, "SharePoint link": 40, "Notes": 30}
    sheet(wb, "Videos", headers, rows, widths, status_col=1, decision_col=2)

    status_of = {v.video_uuid: m for v, m in videos}
    title_of = {v.video_uuid: v.title for v, _ in videos}
    p_rows, c_rows = [], []
    for p in result["playlists"]:
        chapters = p["chapters"]
        in_sp = sum(1 for c in chapters if c.get("status") == ci.IN_SP)
        to_move = sum(1 for c in chapters if c.get("status") == ci.MIGRATE)
        p_rows.append([p["title"], link(p["uuid"]), p["uuid"], p["type"],
                       "yes" if p["public"] else "no",
                       len(chapters), in_sp, len(chapters) - in_sp - to_move, to_move])
        for n, c in enumerate(chapters, start=1):
            m = status_of.get(c["video_uuid"])
            c_rows.append([p["title"], link(p["uuid"]), n, c["chapter"], c["video_uuid"],
                           title_of.get(c["video_uuid"]), c.get("status"),
                           (m.best or {}).get("file") if m else None])
    sheet(wb, "Playlists", ["Playlist", "Open in Consensus", "Playlist UUID", "Type", "Public",
                            "Chapters", "In SharePoint", "To check", "To migrate"],
          p_rows, {"Playlist": 55, "Open in Consensus": 30, "Playlist UUID": 38})
    sheet(wb, "Playlist chapters", ["Playlist", "Open in Consensus", "#", "Chapter",
                                    "Video UUID", "Video", "Status", "SharePoint file"],
          c_rows, {"Playlist": 45, "Open in Consensus": 30, "Chapter": 45, "Video UUID": 38,
                   "Video": 45, "Status": 30, "SharePoint file": 48}, status_col=7)

    shared = [[k, len(t), "; ".join(t)] for k, t in result["shared_files"].items()]
    sheet(wb, "Shared SharePoint files", ["SharePoint item id", "Consensus videos",
                                          "Which"], shared,
          {"SharePoint item id": 38, "Which": 120})
    sheet(wb, "Demos without video", ["Title", "UUID", "Public", "Created"],
          [[d["title"], d["uuid"], "yes" if d["public"] else "no", (d["created_at"] or "")[:10]]
           for d in result["no_video"]], {"Title": 60, "UUID": 38})
    sheet(wb, "Unpublished (not compared)", ["Title", "UUID", "Type", "Draft", "Archived",
                                             "Created"],
          [[d.get("title"), d.get("uuid"), d.get("type"), d.get("draft"), d.get("isArchived"),
            (d.get("createdAt") or "")[:10]] for d in unpublished],
          {"Title": 60, "UUID": 38})

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out", default=os.path.join(
        settings.data_dir, "reports", f"consensus-inventory-{date.today()}.xlsx"))
    ap.add_argument("--details", help="a saved demosDetails JSON instead of fetching")
    args = ap.parse_args()

    fetched_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    if args.details:
        with open(args.details, encoding="utf-8") as fh:
            details = json.load(fh)
        v1_all = []
        print(f"  {len(details)} demos from {args.details}")
    else:
        print("Reading Consensus (read-only)...")
        details = fetch_details()
        print(f"  {len(details)} published demos with video details")
        v1_all = fetch_all_v1()
        print(f"  {len(v1_all)} demos in all states")

    dv, dv_age = mirror("demo_video")
    dc, dc_age = mirror("sharepoint")
    print(f"SharePoint mirror: {len(dv)} Demo Video ({dv_age}), {len(dc)} Demo Catalog ({dc_age})")
    languages = {d.get("uuid"): (d.get("language") or {}).get("code")
                 for d in v1_all if isinstance(d.get("language"), dict)}
    result = ci.build(details, {"Demo Video": dv, "Demo Catalog": dc}, languages)
    result["details"] = details
    write(args.out, result, v1_all, {"demo_video": dv_age, "sharepoint": dc_age},
          fetched_at)
    for status in ci.STATUS_ORDER:
        print(f"  {status}: {sum(1 for _, m in result['videos'] if m.status == status)}")
    print(f"Written: {os.path.abspath(args.out)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
